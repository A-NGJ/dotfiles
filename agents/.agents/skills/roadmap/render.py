# /// script
# requires-python = ">=3.9"
# dependencies = ["python-pptx>=1.0"]
# ///
"""Render a roadmap data file into Markdown, HTML, and PowerPoint.

Usage:
  uv run render.py roadmap.json --out DIR [--format md,html,pptx] [--template corp.pptx] [--layout band|plain]
  uv run render.py --make-template out.pptx [--colors accent1=RRGGBB,dk1=RRGGBB,...] [--font "Typeface"]

--layout overrides the data file's "layout" key (default band). The PowerPoint render prints one
stderr line per lane label or bar whose text won't fit its box; the exit code stays 0.
--make-template writes a minimal 16:9 template: the given theme colour slots (dk1, lt1, dk2, lt2,
accent1..accent6) and theme typeface. Slots left out keep the python-pptx default theme.
"""

from __future__ import annotations

import argparse
import base64
import datetime as dt
import html
import json
import mimetypes
import re
import sys
from pathlib import Path

WEEK = dt.timedelta(weeks=1)
LAYOUTS = ("band", "plain")
SLOTS = ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6")
# python-pptx's default theme, used when no template is given or its theme can't be read.
DEFAULT_THEME = {
    "dk1": "000000", "lt1": "FFFFFF", "dk2": "1F497D", "lt2": "EEECE1", "accent1": "4F81BD",
    "accent2": "C0504D", "accent3": "9BBB59", "accent4": "8064A2", "accent5": "4BACC6", "accent6": "F79646",
}


# --- data ------------------------------------------------------------------------------


class Roadmap:
    def __init__(self, data: dict):
        self.title = data.get("title", "Roadmap")
        self.as_of = _date(data.get("as_of") or dt.date.today().isoformat(), "as_of")
        window = data.get("window") or {}
        start = _date(window.get("start") or (self.as_of - 2 * WEEK).isoformat(), "window.start")
        self.start = start - dt.timedelta(days=start.weekday())  # weeks start on Monday
        self.weeks = [self.start + i * WEEK for i in range(int(window.get("weeks", 12)))]
        self.end = self.weeks[-1] + dt.timedelta(days=6)
        self.milestones = []
        for i, m in enumerate(data.get("milestones", [])):
            where = f"milestones[{i}]"
            for key in ("theme", "name", "start", "finish"):
                if not m.get(key):
                    sys.exit(f"{where}: missing {key}")
            ms = {
                "theme": m["theme"],
                "name": m["name"],
                "start": _date(m["start"], f"{where}.start"),
                "finish": _date(m["finish"], f"{where}.finish"),
                "start_estimated": bool(m.get("start_estimated")),
                "finish_estimated": bool(m.get("finish_estimated")),
            }
            if ms["finish"] < ms["start"]:
                sys.exit(f"{where}: finish is before start")
            self.milestones.append(ms)
        self.milestones.sort(key=lambda m: (m["start"], m["finish"]))
        self.unscheduled = data.get("unscheduled", [])
        self.themes = list(dict.fromkeys(m["theme"] for m in self.milestones))
        self.visible = [m for m in self.milestones if m["finish"] >= self.start and m["start"] <= self.end]
        self.outside = [m for m in self.milestones if m not in self.visible]
        self.layout = data.get("layout", "band")
        if self.layout not in LAYOUTS:
            sys.exit(f"layout: expected one of {', '.join(LAYOUTS)}, got {self.layout!r}")
        self.logo = data.get("logo")  # None: empty slot; False: nothing; str: image path

    def col(self, day: dt.date) -> int:
        """Index of the week column holding `day`, clamped to the window."""
        return min(max((day - self.start).days // 7, 0), len(self.weeks) - 1)

    def span(self, m) -> tuple[int, int]:
        return self.col(m["start"]), self.col(m["finish"])

    def rows(self, theme: str) -> list[list[dict]]:
        """Pack a theme's visible milestones into rows so overlapping ones stack."""
        rows: list[list[dict]] = []
        for m in (m for m in self.visible if m["theme"] == theme):
            for row in rows:
                if self.span(row[-1])[1] < self.span(m)[0]:
                    row.append(m)
                    break
            else:
                rows.append([m])
        return rows

    def lanes(self):
        return [(t, self.rows(t)) for t in self.themes if self.rows(t)]

    def today_offset(self) -> float | None:
        """Position of as_of across the window as a fraction, or None when outside it."""
        if not self.start <= self.as_of <= self.end:
            return None
        return (self.as_of - self.start).days / (7 * len(self.weeks))

    def has_estimates(self) -> bool:
        return any(estimated(m) for m in self.milestones)

    def legend(self) -> list[bool]:
        """Bar states drawn on the timeline, as `estimated` flags: estimated first, then firm."""
        states = {estimated(m) for m in self.visible}
        return [s for s in (True, False) if s in states]


def _date(value: str, where: str) -> dt.date:
    try:
        return dt.date.fromisoformat(value)
    except (TypeError, ValueError):
        sys.exit(f"{where}: expected YYYY-MM-DD, got {value!r}")


def short(day: dt.date) -> str:
    return f"{day.day} {day:%b}"


def dates(m) -> str:
    s = ("~" if m["start_estimated"] else "") + short(m["start"])
    f = ("~" if m["finish_estimated"] else "") + short(m["finish"])
    return f"{s} – {f}"


def estimated(m) -> bool:
    return m["start_estimated"] or m["finish_estimated"]


def as_of_line(r: Roadmap) -> str:
    return f"As of {short(r.as_of)} {r.as_of.year}. Weeks start on Monday. ~ marks an estimated date."


def subtitle(r: Roadmap) -> str:
    """Slide and HTML subtitle; bar styles are left to the legend."""
    parts = [f"As of {short(r.as_of)} {r.as_of.year}", "weeks start on Monday"]
    return " · ".join(parts + (["~ marks an estimated date"] if r.has_estimates() else []))


def unscheduled_lines(r: Roadmap) -> list[str]:
    return [f"{u.get('name', '?')}" + (f" ({u['theme']})" if u.get("theme") else "") for u in r.unscheduled]


LEGEND_LABEL = {True: "Estimated date", False: "Firm dates"}
LOGO_SLOT_TEXT = ("Logo", "(paste official file)")


# --- theme colours ---------------------------------------------------------------------


def theme_colors(prs=None) -> dict[str, str]:
    """RGB hex per theme slot, read from the presentation's first master theme."""
    colors = dict(DEFAULT_THEME)
    if prs is None:
        return colors
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    try:
        theme = etree.fromstring(prs.slide_masters[0].part.part_related_by(RT.THEME).blob)
    except (KeyError, IndexError, etree.XMLSyntaxError):
        return colors
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    for slot in SLOTS:
        el = theme.find(f".//a:clrScheme/a:{slot}/*", ns)
        if el is not None:
            value = el.get("val") if el.tag.endswith("srgbClr") else el.get("lastClr")
            if value and re.fullmatch(r"[0-9A-Fa-f]{6}", value):
                colors[slot] = value.upper()
    return colors


def _luminance(hex_: str) -> float:
    def channel(c: float) -> float:
        return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4

    r, g, b = (channel(int(hex_[i:i + 2], 16) / 255) for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def contrast(a: str, b: str) -> float:
    """WCAG contrast ratio between two RGB hex colours."""
    la, lb = sorted((_luminance(a), _luminance(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def roles(c: dict[str, str]) -> dict[str, str]:
    """Theme slot for each drawing role, chosen by contrast against the fill it sits on."""
    on_accent = max(("lt1", "dk1"), key=lambda s: contrast(c[s], c["accent1"]))
    return {
        "on_accent": on_accent,  # text on an Accent 1 fill
        "accent_ink": "accent1" if contrast(c["accent1"], c["lt1"]) >= 4.5 else "dk1",  # text on white
        "stripe": "accent2" if contrast(c["accent2"], c["accent1"]) >= 3 else on_accent,
    }


def mix(fg: str, bg: str, amount: float) -> str:
    """`fg` at `amount` opacity over `bg`, as RGB hex."""
    return "".join(
        f"{round(int(fg[i:i + 2], 16) * amount + int(bg[i:i + 2], 16) * (1 - amount)):02X}" for i in (0, 2, 4)
    )


# Tints are Accent 1 (or Light 1 over Accent 1) at a fixed opacity.
GRID, BAND, STRONG, PLAIN_LANE, LABEL_RULE = 0.11, 0.04, 0.32, 0.08, 0.25


# --- text fitting ----------------------------------------------------------------------

# Conservative average advances for a bold theme font, in em. PowerPoint wraps wider than the
# QuickLook preview, so these are calibrated against a PowerPoint PDF export, not the thumbnail.
CHAR_EM, SPACE_EM, LINE_EM = 0.62, 0.3, 1.25


def wrap(value: str, width: float, size: float) -> tuple[int, bool]:
    """Greedy word wrap: (line count, whether every word fits the width on its own)."""
    lines, used, words_fit = 0, 0.0, True
    for word in value.split():
        w = len(word) * size * CHAR_EM
        words_fit &= w <= width
        if lines == 0 or used + size * SPACE_EM + w > width:
            lines, used = lines + 1, w
        else:
            used += size * SPACE_EM + w
    return lines, words_fit


def fits(paragraphs: list[tuple[str, float]], width: float, height: float) -> bool:
    total, ok = 0.0, True
    for value, size in paragraphs:
        n, words_fit = wrap(value, width, size)
        total, ok = total + n * size * LINE_EM, ok and words_fit
    return ok and total <= height


# --- Markdown --------------------------------------------------------------------------


def _md_escape(text: str) -> str:
    return text.replace("|", "\\|")


def render_md(r: Roadmap) -> str:
    out = [f"# {r.title}", "", f"_{as_of_line(r)}_", ""]
    out.append("| Theme | " + " | ".join(short(w) for w in r.weeks) + " |")
    out.append("|---|" + "---|" * len(r.weeks))
    for theme, rows in r.lanes():
        for i, row in enumerate(rows):
            cells = [""] * len(r.weeks)
            for m in row:
                a, b = r.span(m)
                cells[a] = f"**{_md_escape(m['name'])}**"
                for j in range(a + 1, b + 1):
                    cells[j] = "▬"
            out.append(f"| {_md_escape(theme) if i == 0 else ''} | " + " | ".join(cells) + " |")
    out += ["", "## Milestones", ""]
    for theme, _ in r.lanes():
        out += [f"**{theme}**", ""]
        out += [f"- {m['name']}: {dates(m)}" for m in r.visible if m["theme"] == theme]
        out.append("")
    if r.outside:
        out += ["## Outside this window", ""]
        out += [f"- {m['name']} ({m['theme']}): {dates(m)}" for m in r.outside]
        out.append("")
    if r.unscheduled:
        out += ["## Unscheduled", ""]
        out += [f"- {line}" for line in unscheduled_lines(r)]
        out.append("")
    return "\n".join(out)


# --- HTML ------------------------------------------------------------------------------

CSS = """
:root{--grid:color-mix(in srgb,var(--accent) 11%,var(--paper));--band:color-mix(in srgb,var(--accent) 4%,var(--paper));
  --strong:color-mix(in srgb,var(--accent) 32%,var(--paper));--muted:color-mix(in srgb,var(--ink) 65%,var(--paper))}
*{box-sizing:border-box}
body{font:14px/1.3 system-ui,-apple-system,"Segoe UI",sans-serif;margin:0;color:var(--ink);background:var(--paper)}
header{position:relative;padding:26px 48px 22px}
h1{margin:0 0 6px;font-size:30px;font-weight:400;padding-right:150px} h2{font-size:16px;margin:24px 0 8px}
.sub{font-size:14px;color:var(--muted)}
.band header{background:var(--accent);color:var(--on-accent)} .band .sub{color:inherit}
.band header::after{content:"";position:absolute;left:48px;bottom:0;width:64px;height:4px;background:var(--stripe)}
.logo{position:absolute;right:48px;top:34px;width:120px;height:40px;display:grid;place-items:center;text-align:center;
  font-size:10px;line-height:1.2;border:1.34px dashed currentColor}
.logo.img{border:0} .logo img{max-width:100%;max-height:100%}
main{padding:24px 48px 16px}
.sl{position:relative;display:grid;grid-template-columns:180px repeat(var(--n),1fr)}
.sl-h{font-size:12px;font-weight:600;color:var(--accent-ink);padding:0 0 6px 6px}
.sl-lane{grid-column:1;font-weight:600;font-size:15px;padding:0 10px 0 14px;display:flex;align-items:center;
  background:var(--lane-bg);color:var(--lane-ink);border-top:1.34px solid var(--lane-rule)}
.sl-lane.first{border-top-color:transparent}
.sl-row{display:grid;grid-template-columns:repeat(var(--n),1fr);align-items:center;min-height:62px;padding:5px 0;
  border-top:1.34px solid var(--grid)}
.sl-row.lane-top{border-top-color:var(--strong)} .sl-row.alt{background:var(--band)}
.cols{grid-column:2 / -1;display:grid;grid-template-columns:repeat(var(--n),1fr);z-index:1;pointer-events:none}
.cols i{border-left:1.34px solid var(--grid)} .cols i:last-child{border-right:1.34px solid var(--grid)}
.sl-row.last{border-bottom:1.34px solid var(--strong)}
.band .sl-lane.first+.sl-row.lane-top{border-top-color:var(--grid)}
.bar{position:relative;z-index:3;margin:0 3px;padding:3px 9px;min-height:52px;border-radius:6px;display:flex;
  flex-direction:column;justify-content:center;background:var(--accent);color:var(--on-accent)}
.bar b{font-size:12.5px;line-height:1.2} .bar span{font-size:10.5px}
.bar.est{background:var(--paper);color:var(--accent-ink);border:1.5px dashed var(--accent)}
.today{position:absolute;top:var(--head);bottom:0;width:2px;margin-left:-1px;background:var(--accent);z-index:2}
.today i{position:absolute;top:calc(100% + 8px);left:50%;transform:translateX(-50%);font:600 10px/1 system-ui,sans-serif;
  padding:3px 7px;border-radius:9px;background:var(--accent);color:var(--on-accent);font-style:normal}
.legend{display:flex;gap:18px;font-size:11px;margin-top:40px;color:var(--accent-ink)}
.legend s{display:inline-block;width:22px;height:12px;border-radius:3px;margin-right:6px;vertical-align:-2px;
  background:var(--accent)}
.legend s.est{background:var(--paper);border:1.5px dashed var(--accent)}
@media print{.sl{break-inside:avoid}}
"""


def _logo_html(r: Roadmap, logo: Path | None) -> str:
    if r.logo is False:
        return ""
    if logo is not None and logo.is_file():
        mime = mimetypes.guess_type(logo.name)[0] or "image/png"
        data = base64.b64encode(logo.read_bytes()).decode()
        return f'<div class="logo img"><img alt="Logo" src="data:{mime};base64,{data}"></div>'
    return f'<div class="logo">{LOGO_SLOT_TEXT[0]}<br>{LOGO_SLOT_TEXT[1]}</div>'


def render_html(r: Roadmap, colors: dict[str, str], logo: Path | None = None) -> str:
    e = html.escape
    n = len(r.weeks)
    role = roles(colors)
    band = r.layout == "band"
    css_vars = {
        "--accent": colors["accent1"], "--ink": colors["dk1"], "--paper": colors["lt1"],
        "--on-accent": colors[role["on_accent"]], "--accent-ink": colors[role["accent_ink"]],
        "--stripe": colors[role["stripe"]],
        "--lane-bg": colors["accent1"] if band else mix(colors["accent1"], colors["lt1"], PLAIN_LANE),
        "--lane-ink": colors[role["on_accent"]] if band else colors["dk1"],
        "--lane-rule": mix(colors["lt1"], colors["accent1"], LABEL_RULE) if band
        else mix(colors["accent1"], colors["lt1"], STRONG),
    }
    root = ";".join(f"{k}:#{v}" for k, v in css_vars.items())
    grid = [f'<div class="sl" style="--n:{n};--head:24px">', '<div class="sl-h"></div>']
    grid += [f'<div class="sl-h">{short(w)}</div>' for w in r.weeks]
    lanes = r.lanes()
    top = 2  # grid row of the current lane; explicit rows keep a multi-row lane label spanning its rows
    for k, (theme, rows) in enumerate(lanes):
        first = " first" if k == 0 else ""
        grid.append(f'<div class="sl-lane{first}" style="grid-row:{top} / span {len(rows)}">{e(theme)}</div>')
        for j, row in enumerate(rows):
            cls = ["sl-row"] + (["lane-top"] if j == 0 else []) + (["alt"] if k % 2 else [])
            cls += ["last"] if k == len(lanes) - 1 and j == len(rows) - 1 else []
            grid.append(f'<div class="{" ".join(cls)}" style="grid-row:{top + j};grid-column:2 / span {n}">')
            for m in row:
                a, b = r.span(m)
                bar = "bar est" if estimated(m) else "bar"
                grid.append(
                    f'<div class="{bar}" style="grid-column:{a + 1} / {b + 2}" title="{e(m["name"])}: {dates(m)}">'
                    f"<b>{e(m['name'])}</b><span>{dates(m)}</span></div>"
                )
            grid.append("</div>")
        top += len(rows)
    body_rows = sum(len(rows) for _, rows in lanes)
    grid.append(f'<div class="cols" style="grid-row:2 / span {body_rows}">' + "<i></i>" * n + "</div>")
    offset = r.today_offset()
    if offset is not None:
        grid.append(f'<div class="today" title="{short(r.as_of)}" '
                    f'style="left:calc(180px + (100% - 180px) * {offset:.4f})"><i>Today</i></div>')
    grid.append("</div>")
    legend = ['<div class="legend">']
    legend += [f'<span><s class="{"est" if s else ""}"></s>{LEGEND_LABEL[s]}</span>' for s in r.legend()]
    legend.append("</div>")
    extra = []
    if r.outside:
        extra.append("<h2>Outside this window</h2><ul>")
        extra += [f"<li>{e(m['name'])} ({e(m['theme'])}): {dates(m)}</li>" for m in r.outside]
        extra.append("</ul>")
    if r.unscheduled:
        extra.append("<h2>Unscheduled</h2><ul>")
        extra += [f"<li>{e(line)}</li>" for line in unscheduled_lines(r)]
        extra.append("</ul>")
    return (
        f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{e(r.title)}</title>'
        f"<style>:root{{{root}}}{CSS}</style></head><body class=\"{r.layout}\"><header><h1>{e(r.title)}</h1>"
        f'<div class="sub">{e(subtitle(r))}</div>{_logo_html(r, logo)}</header><main>'
        + "\n".join(grid + legend + extra)
        + "</main></body></html>\n"
    )


# --- PowerPoint ------------------------------------------------------------------------

# Geometry in a 1280x720 px design space scaled to the slide width (1 px = 9525 EMU on 13.333 in).
LEFT, RIGHT, LANE_W, HEAD_H, HEADER_H, CHART_TOP = 48, 1232, 180, 22, 128, 168
RULE = 1.34  # px; 1 pt, the thinnest rule every viewer draws
NAME_PX, DATES_PX, LANE_PX = (12.5, 11.5), 10.5, 15  # a bar name shrinks one step before it warns
BAR_PAD = (9, 3, 9, 3)  # left, top, right, bottom text margins inside a bar
LANE_PAD = (14, 0, 10, 0)
LOGO_BOX = (1112, 34, 120, 40)
LOGO_TINT = 0.12  # slot fill, so the slot shows where a viewer drops the dashed line


def render_pptx(r: Roadmap, path: Path, template: Path | None, logo: Path | None = None) -> list[str]:
    """Write the slides; return one overflow warning per lane label or bar that won't fit."""
    from pptx import Presentation
    from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_THEME_COLOR
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt

    prs = Presentation(str(template)) if template else Presentation()
    if template:  # keep the masters and layouts, drop the template's own slides
        ids = prs.slides._sldIdLst
        for sld in list(ids):
            prs.part.drop_rel(sld.rId)
            ids.remove(sld)
    else:
        prs.slide_width, prs.slide_height = Emu(1280 * 9525), Emu(720 * 9525)
    colors = theme_colors(prs)
    role = roles(colors)
    band = r.layout == "band"
    unit = prs.slide_width / 1280
    warnings: list[str] = []

    def px(v: float) -> Emu:
        return Emu(int(round(v * unit)))

    def fpt(v: float) -> Pt:  # design px to points, at 1280 px = 13.333 in
        return Pt(v * 0.75)

    def paint(fmt, slot: str, alpha: float | None = None, brightness: float = 0.0):
        """Theme colour, optionally lightened (brightness) or made translucent (alpha)."""
        # Slides reference dk1/lt1/dk2/lt2 through the master's colour map (tx1/bg1/tx2/bg2).
        fmt.theme_color = getattr(MSO_THEME_COLOR, {"dk1": "TEXT_1", "lt1": "BACKGROUND_1", "dk2": "TEXT_2",
                                                     "lt2": "BACKGROUND_2"}.get(slot, slot.upper().replace("ACCENT", "ACCENT_")))
        if brightness:
            fmt.brightness = brightness
        if alpha is not None:
            el = fmt._color._xClr
            for old in el.findall(qn("a:alpha")):
                el.remove(old)
            el.append(el.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))}))

    def shape(slide, x, y, w, h, fill=None, alpha=None, kind=MSO_SHAPE.RECTANGLE):
        s = slide.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
        s.shadow.inherit = False
        if fill is None:
            s.fill.background()
        else:
            s.fill.solid()
            paint(s.fill.fore_color, fill, alpha)
        s.line.fill.background()
        return s

    def outline(s, slot, width=1.5, dashed=True):
        # A theme reference, like every other colour. QuickLook draws theme-coloured lines black
        # (PowerPoint draws them right), so a line that must show on a dark fill also gets a tint.
        paint(s.line.color, slot)
        s.line.width = fpt(max(width, RULE))
        if dashed:
            s.line.dash_style = MSO_LINE_DASH_STYLE.DASH

    def write(target, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, margins=(0, 0, 0, 0)):
        """paras: (text, size_px, bold, ink) per paragraph. ink is a slot or (slot, brightness);
        a size or ink of None inherits from the placeholder. Fonts always come from the theme."""
        tf = target.text_frame
        tf.word_wrap, tf.vertical_anchor = True, anchor
        tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (px(m) for m in margins)
        tf.clear()
        for i, (value, size, bold, ink) in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            run = p.add_run()
            run.text = value
            run.font.bold = bold
            if size is not None:
                run.font.size = fpt(size)
            if ink is not None:
                slot, light = ink if isinstance(ink, tuple) else (ink, 0.0)
                paint(run.font.color, slot, brightness=light)
        return target

    def text(slide, x, y, w, h, paras, **kw):
        return write(slide.shapes.add_textbox(px(x), px(y), px(w), px(h)), paras, **kw)

    def to_back(slide, s):
        tree = slide.shapes._spTree
        tree.remove(s._element)
        tree.insert(2, s._element)

    def layout():
        by_name = {l.name.lower(): l for l in prs.slide_layouts}
        return by_name.get("title only") or min(prs.slide_layouts, key=lambda l: len(l.placeholders))

    def header(title: str):
        slide = prs.slides.add_slide(layout())
        if band:
            to_back(slide, shape(slide, 0, 0, 1280, HEADER_H, "accent1"))
            shape(slide, LEFT, HEADER_H - 4, 64, 4, role["stripe"])
        title_w = (LOGO_BOX[0] - 16 if r.logo is not False else RIGHT) - LEFT
        t = slide.shapes.title or slide.shapes.add_textbox(0, 0, 0, 0)
        t.left, t.top, t.width, t.height = px(LEFT), px(26), px(title_w), px(52)
        if band:
            line = (title, 40, False, role["on_accent"])
        elif template:  # plain on a template: its own title size and colour
            line = (title, None, None, None)
        else:
            line = (title, 37, False, "dk1")
        write(t, [line], anchor=MSO_ANCHOR.TOP)
        sub_ink = role["on_accent"] if band else ("dk1", 0.35)
        text(slide, LEFT, 82, title_w, 22, [(subtitle(r), 14, False, sub_ink)])
        draw_logo(slide)
        return slide

    def draw_logo(slide):
        if r.logo is False:
            return
        x, y, w, h = LOGO_BOX
        if logo is not None and logo.is_file():
            pic = slide.shapes.add_picture(str(logo), px(x), px(y))
            scale = min(px(w) / pic.width, px(h) / pic.height)
            pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
            pic.left, pic.top = px(x + w) - pic.width, px(y)
            return
        ink = role["on_accent"] if band else "accent1"
        slot = shape(slide, x, y, w, h, ink, LOGO_TINT)
        outline(slot, ink, RULE)
        write(slot, [(line, 10, False, ink) for line in LOGO_SLOT_TEXT], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)

    # --- timeline slide ---------------------------------------------------------------
    slide = header(r.title)
    grid_x = LEFT + LANE_W
    grid_w = RIGHT - grid_x
    n = len(r.weeks)
    col_w = grid_w / n
    lanes = r.lanes()
    total_rows = sum(len(rows) for _, rows in lanes) or 1
    row_h = min(62, 470 // total_rows)
    body_top = CHART_TOP + HEAD_H
    body_bottom = body_top + row_h * total_rows

    for i, week in enumerate(r.weeks):
        text(slide, grid_x + col_w * i + 6, CHART_TOP, col_w - 6, HEAD_H, [(short(week), 12, True, role["accent_ink"])])

    y = body_top
    spans = []
    for k, (theme, rows) in enumerate(lanes):
        h = row_h * len(rows)
        if k % 2:
            shape(slide, grid_x, y, grid_w, h, "accent1", BAND)
        label = shape(slide, LEFT, y, LANE_W, h, "accent1", None if band else PLAIN_LANE)
        write(label, [(theme, LANE_PX, True, role["on_accent"] if band else "dk1")],
              anchor=MSO_ANCHOR.MIDDLE, margins=LANE_PAD)
        if not fits([(theme, LANE_PX)], LANE_W - LANE_PAD[0] - LANE_PAD[2], h):
            warnings.append(f'{r.layout}: lane label "{theme}" overflows its {LANE_W}px label cell')
        spans.append((y, h, rows))
        y += h

    for ly, h, rows in spans:  # rules between rows inside a lane
        for j in range(1, len(rows)):
            shape(slide, grid_x, ly + row_h * j - RULE / 2, grid_w, RULE, "accent1", GRID)
    for i in range(n + 1):  # a rule at every week boundary, including the right edge
        x = grid_x + col_w * i - (RULE if i == n else 0)
        shape(slide, x, body_top, RULE, body_bottom - body_top, "accent1", GRID)
    for k, (ly, h, _) in enumerate(spans):  # stronger rules between lanes, across the label column
        edge = ly + h - (RULE if k == len(spans) - 1 else RULE / 2)
        shape(slide, grid_x, edge, grid_w, RULE, "accent1", STRONG)
        if k < len(spans) - 1:
            if band:
                shape(slide, LEFT, edge, LANE_W, RULE, "lt1", LABEL_RULE)
            else:
                shape(slide, LEFT, edge, LANE_W, RULE, "accent1", STRONG)

    offset = r.today_offset()  # drawn before the bars so it sits behind them
    if offset is not None:
        tx = grid_x + grid_w * offset
        shape(slide, tx - 1, body_top, 2, body_bottom - body_top, "accent1")

    for ly, _, rows in spans:
        for j, row in enumerate(rows):
            for m in row:
                a, b = r.span(m)
                x, w, h = grid_x + col_w * a + 3, col_w * (b - a + 1) - 6, row_h - 10
                firm = not estimated(m)
                bar = shape(slide, x, ly + row_h * j + 5, w, h, "accent1" if firm else "lt1",
                            kind=MSO_SHAPE.ROUNDED_RECTANGLE)
                bar.adjustments[0] = min(0.5, 6 / min(w, h))
                if not firm:
                    outline(bar, "accent1")
                ink = role["on_accent"] if firm else role["accent_ink"]
                inner_w, inner_h = w - BAR_PAD[0] - BAR_PAD[2], h - BAR_PAD[1] - BAR_PAD[3]
                size = next((s for s in NAME_PX if fits([(m["name"], s), (dates(m), DATES_PX)], inner_w, inner_h)), None)
                if size is None:
                    size = NAME_PX[-1]
                    warnings.append(f'{r.layout}: bar "{m["name"]}" ({m["theme"]}) overflows its {b - a + 1}-week bar')
                paras = [(m["name"], size, True, ink), (dates(m), DATES_PX, False, ink)]
                write(bar, paras, anchor=MSO_ANCHOR.MIDDLE, margins=BAR_PAD)

    if offset is not None:
        pill = shape(slide, tx - 21, body_bottom + 8, 42, 15, "accent1", kind=MSO_SHAPE.ROUNDED_RECTANGLE)
        pill.adjustments[0] = 0.5
        write(pill, [("Today", 10, True, role["on_accent"])], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)

    lx = LEFT
    for state in r.legend():
        key = shape(slide, lx, 688, 22, 12, "lt1" if state else "accent1", kind=MSO_SHAPE.ROUNDED_RECTANGLE)
        if state:
            outline(key, "accent1")
        text(slide, lx + 28, 686, 110, 16, [(LEGEND_LABEL[state], 11, False, role["accent_ink"])])
        lx += 140

    # --- not-on-the-timeline slide ------------------------------------------------------
    sections = [
        ("Outside this window", [f"{m['name']} ({m['theme']}): {dates(m)}" for m in r.outside]),
        ("Unscheduled", unscheduled_lines(r)),
    ]
    sections = [(t, lines) for t, lines in sections if lines]
    if sections:
        slide = header(f"{r.title}: not on the timeline")
        paras = []
        for heading, lines in sections:
            paras.append((heading, 20, True, "dk1"))
            paras += [(f"•  {line}", 16, False, "dk1") for line in lines]
        box = text(slide, LEFT, 176, RIGHT - LEFT, 480, paras)
        for p in box.text_frame.paragraphs:
            p.space_after = Pt(8)

    prs.save(str(path))
    return warnings


# --- template builder ------------------------------------------------------------------


def make_template(path: Path, colors: dict[str, str], font: str | None) -> None:
    """Write a minimal 16:9 template: theme colour slots, theme typeface, and a title placed for 16:9."""
    from lxml import etree
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.util import Emu

    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    p = "http://schemas.openxmlformats.org/presentationml/2006/main"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(1280 * 9525), Emu(720 * 9525)
    master = prs.slide_masters[0]
    theme_part = master.part.part_related_by(RT.THEME)
    theme = etree.fromstring(theme_part.blob)
    scheme = theme.find(f".//{{{a}}}clrScheme")
    scheme.set("name", "Custom")
    for name, hex_ in colors.items():
        slot = scheme.find(f"{{{a}}}{name}")
        for child in list(slot):
            slot.remove(child)
        etree.SubElement(slot, f"{{{a}}}srgbClr", val=hex_)
    if font:
        for kind in ("majorFont", "minorFont"):
            theme.find(f".//{{{a}}}{kind}/{{{a}}}latin").set("typeface", font)
    theme_part._blob = etree.tostring(theme, xml_declaration=True, encoding="UTF-8", standalone=True)

    title_style = master._element.find(f".//{{{p}}}titleStyle/{{{a}}}lvl1pPr/{{{a}}}defRPr")
    title_style.set("sz", "2800")
    title_style.set("b", "1")
    for fill in title_style.findall(f"{{{a}}}solidFill"):
        title_style.remove(fill)
    fill = etree.Element(f"{{{a}}}solidFill")
    etree.SubElement(fill, f"{{{a}}}schemeClr", val="accent1")
    title_style.insert(0, fill)
    unit = 9525
    for placeholders in [master.placeholders] + [l.placeholders for l in master.slide_layouts]:
        for ph in placeholders:
            if "TITLE" in str(ph.placeholder_format.type) and "SUB" not in str(ph.placeholder_format.type):
                ph.left, ph.top, ph.width, ph.height = Emu(38 * unit), Emu(29 * unit), Emu(1204 * unit), Emu(77 * unit)
                ph.text_frame.word_wrap = True
    prs.save(str(path))


def parse_colors(spec: str) -> dict[str, str]:
    colors = {}
    for item in filter(None, (s.strip() for s in spec.split(","))):
        slot, _, value = item.partition("=")
        slot, value = slot.strip(), value.strip().lstrip("#").upper()
        if slot not in SLOTS:
            sys.exit(f"--colors: unknown slot {slot!r}; expected one of {', '.join(SLOTS)}")
        if not re.fullmatch(r"[0-9A-F]{6}", value):
            sys.exit(f"--colors: {slot} expects RRGGBB, got {value!r}")
        colors[slot] = value
    return colors


# --- CLI -------------------------------------------------------------------------------


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("data", type=Path, nargs="?", help="roadmap data file (JSON)")
    ap.add_argument("--out", type=Path, help="output directory")
    ap.add_argument("--format", default="md", help="comma-separated: md, html, pptx (default md)")
    ap.add_argument("--template", type=Path, help="corporate .pptx whose masters and theme to use")
    ap.add_argument("--layout", choices=LAYOUTS, help="band or plain; overrides the data file's layout")
    ap.add_argument("--make-template", type=Path, metavar="OUT", help="write a minimal template and exit")
    ap.add_argument("--colors", default="", help="with --make-template: slot=RRGGBB,... (dk1, lt1, dk2, lt2, accent1..6)")
    ap.add_argument("--font", help="with --make-template: theme typeface for headings and body")
    args = ap.parse_args(argv)

    if args.make_template:
        make_template(args.make_template, parse_colors(args.colors), args.font)
        print(args.make_template)
        return 0
    if args.data is None or args.out is None:
        ap.error("data and --out are required unless --make-template is given")
    formats = [f.strip() for f in args.format.split(",") if f.strip()]
    unknown = set(formats) - {"md", "html", "pptx"}
    if unknown:
        ap.error(f"unknown format(s): {', '.join(sorted(unknown))}")
    r = Roadmap(json.loads(args.data.read_text()))
    if not r.milestones:
        sys.exit("no scheduled milestones; nothing to draw")
    if args.layout:
        r.layout = args.layout
    logo = None
    if isinstance(r.logo, str):
        logo = (args.data.parent / r.logo).resolve()
        if not logo.is_file():
            print(f"logo: {logo} not found; drawing the empty logo slot", file=sys.stderr)
    args.out.mkdir(parents=True, exist_ok=True)
    stem = args.data.stem
    written = []
    if "md" in formats:
        written.append(args.out / f"{stem}.md")
        written[-1].write_text(render_md(r))
    if "html" in formats:
        colors = DEFAULT_THEME
        if args.template:
            from pptx import Presentation

            colors = theme_colors(Presentation(str(args.template)))
        written.append(args.out / f"{stem}.html")
        written[-1].write_text(render_html(r, colors, logo))
    if "pptx" in formats:
        written.append(args.out / f"{stem}.pptx")
        for warning in render_pptx(r, written[-1], args.template, logo):
            print(f"overflow: {warning}", file=sys.stderr)
    for path in written:
        print(path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
